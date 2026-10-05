#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import copy
import datetime
import json
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
ACCENT = RGBColor(0xF5, 0x6C, 0x5F)
BODY = RGBColor(0x33, 0x38, 0x3F)
MUTED = RGBColor(0x64, 0x6E, 0x7B)


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def fmt_num(value) -> str:
    return f"{int(value or 0):,}"


def fmt_compact(value) -> str:
    value = int(value or 0)
    if value >= 100_000_000:
        return f"{value / 100_000_000:.2f} 亿"
    if value >= 10_000:
        return f"{value / 10_000:.1f} 万"
    return f"{value:,}"


def fmt_duration(seconds) -> str:
    seconds = int(seconds or 0)
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours:d}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:d}:{secs:02d}"


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def set_run(run, size=10.5, bold=False, color=BODY):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def para(doc, text="", size=10.5, bold=False, color=BODY, style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(str(text))
        set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    return p


def heading(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = color or (ACCENT if level <= 2 else RGBColor(0x21, 0x66, 0xB5))
    return h


def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    set_run(run, size=8.5, bold=bold)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = True
    for idx, value in enumerate(headers):
        set_cell(table.rows[0].cells[idx], value, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell(cells[idx], value)
    if widths:
        for idx, width in enumerate(widths):
            for row in table.rows:
                row.cells[idx].width = Inches(width)
    para(doc)
    return table


def add_image(doc, filename, caption="", width=6.4):
    path = ROOT / "charts" / filename
    if not path.exists():
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, size=8.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)


def parse_quiz():
    path = ROOT / "site" / "data-quiz.js"
    if not path.exists():
        return []
    raw = path.read_text(encoding="utf-8")
    body = raw.split("const QUIZ_DATA =", 1)[-1].rsplit("];", 1)[0] + "]"
    return json.loads(body)


def stat_of(video, key):
    return int(((video.get("_meta") or {}).get("stat") or {}).get(key, 0) or 0)


def fallback_summary(video):
    desc = clean_text(video.get("desc") or (video.get("_meta") or {}).get("desc") or "")
    base = f"标题《{video['title']}》指向{video['type']}内容。"
    if desc:
        base += f"官方简介补充：{desc[:420]}"
    else:
        base += "本视频未提供有效简介，词条依据标题、分区、标签和公开互动数据整理。"
    return base


def fallback_highlights(video):
    tags = " / ".join(video.get("tags", [])[:8]) or video.get("tname") or "暂无标签"
    return [
        f"内容类别：{video['type']}，B站分区为 {video.get('tname') or '未标注'}。",
        f"主要标签：{tags}。",
        f"公开互动：播放 {fmt_compact(stat_of(video, 'view'))}，点赞 {fmt_compact(stat_of(video, 'like'))}，弹幕 {fmt_compact(stat_of(video, 'danmaku'))}。",
        f"时长 {fmt_duration(video.get('duration') or 0)}，投稿间隔 {video.get('gap_note') or '首期投稿'}。",
    ]


def video_info_rows(video):
    return [
        ["编号 / BV号", f"第 {video['no']} 部 / {video['bvid']}"],
        ["投稿时间", video.get("pubdate_iso") or video["date"]],
        ["账号", video["account"]],
        ["类型 / 分区", f"{video['type']} / {video.get('tname') or '—'}"],
        ["时长", fmt_duration(video["duration"])],
        ["播放 / 点赞", f"{fmt_compact(stat_of(video, 'view'))} / {fmt_compact(stat_of(video, 'like'))}"],
        ["弹幕 / 评论", f"{fmt_compact(stat_of(video, 'danmaku'))} / {fmt_compact(stat_of(video, 'reply'))}"],
        ["投币 / 收藏", f"{fmt_compact(stat_of(video, 'coin'))} / {fmt_compact(stat_of(video, 'favorite'))}"],
        ["字幕", f"{video['sub_lines']:,} 行 / {video['sub_chars']:,} 字符"],
        ["距上次投稿", video.get("gap_note") or "首期投稿"],
        ["视频链接", f"https://www.bilibili.com/video/{video['bvid']}"],
    ]


def add_comments(doc, video, comments, count=2):
    rows = comments.get(video["bvid"], {}).get("comments") or []
    if not rows:
        return
    heading(doc, "热门评论", level=4)
    for item in rows[:count]:
        message = clean_text(item.get("message", ""))
        reply = ""
        if item.get("replies"):
            first = item["replies"][0]
            reply = " 热评回复：" + clean_text(first.get("message", ""))
        para(doc, f"{item.get('name', '匿名')}（{fmt_num(item.get('like', 0))} 赞）：{message}{reply}", style="List Bullet")


def add_danmaku(doc, video, insights):
    item = insights.get("videos", {}).get(video["bvid"])
    if not item:
        return
    heading(doc, "弹幕特征", level=4)
    phrases = "、".join(x["content"] for x in item.get("top_phrases", [])[:6]) or "—"
    peak = item.get("peaks", [{}])[0]
    peak_time = int(peak.get("t", 0) or 0)
    samples = " / ".join(clean_text(x)[:44] for x in peak.get("samples", [])[:3]) or "—"
    para(doc, f"高频弹幕：{phrases}", style="List Bullet")
    para(doc, f"峰值窗口：{peak_time // 60:02d}:{peak_time % 60:02d}（{peak.get('count', 0)} 条）｜代表样本：{samples}", style="List Bullet")


def add_video_entry(doc, video, comments, insights):
    heading(doc, f"第{video['no']}部  {video['title']}", level=3)
    add_table(doc, ["项目", "内容"], video_info_rows(video), widths=[1.35, 5.05])
    tags = " / ".join(video.get("tags", [])[:16])
    if tags:
        para(doc, f"标签：{tags}")

    analysis = video.get("analysis")
    if analysis:
        para(doc, "关键话题：" + (analysis.get("topic_str") or "—"), bold=True)
        if analysis.get("catch_str"):
            para(doc, "标志性口头禅：" + analysis["catch_str"], bold=True)
        para(doc, "内容摘要：", bold=True)
        para(doc, analysis.get("summary") or "—")
        if analysis.get("highlights"):
            para(doc, "内容亮点：", bold=True)
            for line in analysis["highlights"]:
                para(doc, line, style="List Bullet")
        if analysis.get("atmo"):
            para(doc, "整体氛围：" + "\n".join(analysis["atmo"]), color=MUTED)
        para(doc, "语言特征：" + (analysis.get("lang") or "—"), color=MUTED)
        if analysis.get("merged"):
            para(doc, f"字幕全文：共 {video['sub_lines']:,} 行，合并为 {len(analysis['merged'])} 段。", bold=True)
            for merged in analysis["merged"]:
                p = para(doc, merged, size=9)
                p.paragraph_format.space_after = Pt(1)
    else:
        para(doc, "公开数据补录：", bold=True)
        para(doc, fallback_summary(video))
        para(doc, "词条要点：", bold=True)
        for line in fallback_highlights(video):
            para(doc, line, style="List Bullet")

    add_comments(doc, video, comments)
    add_danmaku(doc, video, insights)
    para(doc, f"— 第{video['no']}部完 —", size=8.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)


class DocPatcher:
    """词条增量补丁引擎：克隆既有模板段落/表格，追加词条并同步统计（与 Warma 百科同构）。"""

    def __init__(self, path):
        self.path = Path(path)
        self.doc = Document(str(path))
        self.body = self.doc.element.body
        self._collect_templates()

    # ---------- 模板 ----------
    def _collect_templates(self):
        tpls = {}
        for p in self.doc.paragraphs:
            t = p.text.strip()
            st = p.style.name
            if "tpl_h3" not in tpls and st == "Heading 3" and re.match(r"^第\d+部", t):
                tpls["tpl_h3"] = p._p
            elif "tpl_label" not in tpls and t == "内容摘要：":
                tpls["tpl_label"] = p._p
            elif "tpl_body" not in tpls and st == "Normal" and len(t) > 40:
                for r in p.runs:
                    if r.font.size and r.font.size.pt <= 10.6 and not r.bold:
                        tpls["tpl_body"] = p._p
                        break
            elif "tpl_meta" not in tpls and t.startswith("整体氛围："):
                tpls["tpl_meta"] = p._p
            elif "tpl_lang" not in tpls and t.startswith("语言特征："):
                tpls["tpl_lang"] = p._p
            elif "tpl_note" not in tpls and re.match(r"^字幕全文：共.*行.*合并为.*段", t):
                tpls["tpl_note"] = p._p
            elif "tpl_end" not in tpls and re.match(r"^— 第\d+部完 —$", t):
                tpls["tpl_end"] = p._p
            elif "tpl_list" not in tpls and st == "List Bullet" and t:
                tpls["tpl_list"] = p._p
            elif "tpl_h2year" not in tpls and st == "Heading 2" and re.match(r"^2\.\d+\s*\d{4}年：", t):
                tpls["tpl_h2year"] = p._p
        if "tpl_note" in tpls:
            seen_note = False
            for p in self.doc.paragraphs:
                if p._p is tpls["tpl_note"]:
                    seen_note = True
                    continue
                if seen_note:
                    if p.style.name == "Normal" and p.text.strip():
                        tpls["tpl_sub"] = p._p
                        break
        self.tpl = {k: copy.deepcopy(v) for k, v in tpls.items()}
        for t in self.doc.tables:
            hdr = [c.text.strip() for c in t.rows[0].cells]
            if "tbl_info" not in self.tpl and hdr[:2] == ["项目", "内容"] and len(t.rows) <= 12:
                self.tpl["tbl_info"] = copy.deepcopy(t._tbl)
            if "tbl_year" not in self.tpl and hdr[:4] == ["编号", "日期", "标题", "类型"]:
                self.tpl["tbl_year"] = copy.deepcopy(t._tbl)
        missing = {"tpl_h3", "tpl_label", "tpl_body", "tpl_end", "tpl_list", "tbl_info", "tbl_year"} - set(self.tpl)
        if missing:
            raise RuntimeError(f"模板定位失败：缺少 {sorted(missing)}")

    # ---------- 基础构件 ----------
    def _mk_para(self, tpl_key, text):
        p = copy.deepcopy(self.tpl[tpl_key])
        for r in p.findall(qn("w:r")):
            p.remove(r)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        r.append(rpr)
        lines = text.split("\n")
        for i, line in enumerate(lines):
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = line
            r.append(t)
            if i < len(lines) - 1:
                r.append(OxmlElement("w:br"))
        p.append(r)
        return p

    def _label_para(self, label, content=""):
        p = copy.deepcopy(self.tpl["tpl_label"])
        runs = p.findall(qn("w:r"))
        rpr = runs[0].find(qn("w:rPr")) if runs else None
        for r in runs:
            p.remove(r)
        r1 = OxmlElement("w:r")
        if rpr is not None:
            r1.append(copy.deepcopy(rpr))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = label
        r1.append(t)
        p.append(r1)
        if content:
            r2 = OxmlElement("w:r")
            if rpr is not None:
                r2.append(copy.deepcopy(rpr))
            t2 = OxmlElement("w:t")
            t2.set(qn("xml:space"), "preserve")
            t2.text = content
            r2.append(t2)
            p.append(r2)
        return p

    def _list_para(self, text):
        p = copy.deepcopy(self.tpl["tpl_list"])
        for r in p.findall(qn("w:r")):
            p.remove(r)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p.append(r)
        return p

    def _tbl_info(self, fields):
        tbl = copy.deepcopy(self.tpl["tbl_info"])
        rows = tbl.findall(qn("w:tr"))
        row_tpl = copy.deepcopy(rows[1])
        for r in rows[1:]:
            tbl.remove(r)
        for k, v in fields:
            tr = copy.deepcopy(row_tpl)
            cells = tr.findall(qn("w:tc"))
            self._cell_set(cells[0], k)
            self._cell_set(cells[1], v)
            tbl.append(tr)
        return tbl

    def _cell_set(self, tc, text):
        ps = tc.findall(qn("w:p"))
        p = ps[0]
        runs = p.findall(qn("w:r"))
        if runs:
            keep = runs[0]
            for r in runs[1:]:
                p.remove(r)
            for t in keep.findall(qn("w:t")):
                keep.remove(t)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            keep.append(t)
        else:
            r = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = text
            r.append(t)
            p.append(r)
        for extra in ps[1:]:
            tc.remove(extra)

    # ---------- 定位 ----------
    def find_h1(self, prefix):
        for p in self.doc.paragraphs:
            if p.style.name == "Heading 1" and p.text.strip().startswith(prefix):
                return p
        return None

    def find_entry_anchor(self):
        """最后一个词条“完”标记段落（附录 H1 之前）。"""
        h1 = self.find_h1("附录")
        if h1 is None:
            raise RuntimeError("未找到附录 H1")
        prev = h1._p.getprevious()
        while prev is not None:
            if prev.tag.endswith("}p"):
                txt = "".join(t.text or "" for t in prev.iter(qn("w:t")))
                if re.match(r"^— 第\d+部完 —$", txt.strip()):
                    return prev
            prev = prev.getprevious()
        raise RuntimeError("未找到插入锚点")

    def find_year_sections(self):
        out = {}
        cur_year = None
        for child in self.body.iterchildren():
            if child.tag.endswith("}p"):
                p = Paragraph(child, self.doc)
                m = re.match(r"^2\.\d+\s*(\d{4})年：(\d+)部投稿$", p.text.strip())
                if p.style.name == "Heading 2" and m:
                    cur_year = int(m.group(1))
                    out[cur_year] = {"h2": p, "count": int(m.group(2)), "table": None}
            elif child.tag.endswith("}tbl") and cur_year is not None:
                t = Table(child, self.doc)
                hdr = [c.text.strip() for c in t.rows[0].cells]
                if hdr[:4] == ["编号", "日期", "标题", "类型"] and out[cur_year]["table"] is None:
                    out[cur_year]["table"] = child
        return out

    # ---------- 词条插入 ----------
    def entry_elements(self, v):
        a = v.get("analysis") or {}
        stat = v.get("stat") or {}

        def s(key):
            return fmt_compact(stat.get(key, 0)) if stat else "—"

        els = []
        els.append(self._mk_para("tpl_h3", f"第{v['no']}部  {v['title']}"))
        els.append(self._tbl_info([
            ("编号 / BV号", f"第 {v['no']} 部 / {v['bvid']}"),
            ("投稿时间", v["date"].replace("-", "/")),
            ("账号", v.get("account") or "—"),
            ("类型 / 分区", f"{v['type']} / {v.get('tname') or '—'}"),
            ("时长", fmt_duration(v.get("duration")) if v.get("duration") else "—"),
            ("播放 / 点赞", f"{s('view')} / {s('like')}"),
            ("弹幕 / 评论", f"{s('danmaku')} / {s('reply')}"),
            ("投币 / 收藏", f"{s('coin')} / {s('favorite')}"),
            ("字幕", f"{v.get('sub_lines', 0):,} 行 / {v.get('sub_chars', 0):,} 字符"),
            ("距上次投稿", v.get("gap_note") or "首期投稿"),
            ("视频链接", f"https://www.bilibili.com/video/{v['bvid']}" if v.get("bvid") else "—"),
        ]))
        els.append(self._label_para("关键话题：", a.get("topic_str") or "—"))
        if a.get("catch_str"):
            els.append(self._label_para("标志性口头禅：", a["catch_str"]))
        els.append(self._label_para("内容摘要："))
        els.append(self._mk_para("tpl_body", a.get("summary") or fallback_summary(v)))
        els.append(self._label_para("内容亮点："))
        for h in a.get("highlights") or fallback_highlights(v):
            els.append(self._list_para(h))
        if a.get("atmo"):
            els.append(self._mk_para("tpl_meta", "整体氛围：" + "\n".join(a["atmo"])))
        if a.get("lang"):
            els.append(self._mk_para("tpl_lang", "语言特征：" + a["lang"]))
        merged = a.get("merged") or []
        if merged:
            els.append(self._mk_para("tpl_note", f"字幕全文：共 {v.get('sub_lines', 0):,} 行，合并为 {len(merged)} 段。"))
            for mtext in merged:
                els.append(self._mk_para("tpl_sub", mtext))
        els.append(self._mk_para("tpl_end", f"— 第{v['no']}部完 —"))
        return els

    def append_entries(self, videos):
        cursor = self.find_entry_anchor()
        for v in videos:
            for el in self.entry_elements(v):
                cursor.addnext(el)
                cursor = el

    # ---------- 年表 ----------
    def _table_row_clone(self, tbl, cells, row_tpl=None):
        rows = tbl.findall(qn("w:tr"))
        tpl = copy.deepcopy(row_tpl) if row_tpl is not None else copy.deepcopy(rows[-1])
        tr = copy.deepcopy(tpl)
        tcs = tr.findall(qn("w:tc"))
        for tc, text in zip(tcs, cells):
            self._cell_set(tc, text)
        tbl.append(tr)
        return tr

    def update_years(self, year_counts):
        secs = self.find_year_sections()
        for year, rows in sorted(year_counts.items()):
            if year in secs and secs[year]["table"] is not None:
                h2 = secs[year]["h2"]
                old_n = secs[year]["count"]
                new_n = old_n + len(rows)
                self._rewrite_text(h2._p, re.sub(r"：\d+部投稿$", f"：{new_n}部投稿", h2.text.strip()))
                tbl = secs[year]["table"]
                for r in rows:
                    self._table_row_clone(tbl, [str(r[0]), r[1], r[2], r[3]])
            else:
                if "tpl_h2year" not in self.tpl or not secs:
                    continue
                idx = max(int(re.match(r"^2\.(\d+)", s["h2"].text.strip()).group(1)) for s in secs.values()) + 1
                h2el = copy.deepcopy(self.tpl["tpl_h2year"])
                self._rewrite_text(h2el, f"2.{idx} {year}年：{len(rows)}部投稿")
                tbl = copy.deepcopy(self.tpl["tbl_year"])
                rows_el = tbl.findall(qn("w:tr"))
                data_tpl = copy.deepcopy(rows_el[1]) if len(rows_el) > 1 else None
                for r in rows_el[1:]:
                    tbl.remove(r)
                for r in rows:
                    self._table_row_clone(tbl, [str(r[0]), r[1], r[2], r[3]], row_tpl=data_tpl)
                h1 = self.find_h1("第三章")
                if h1 is not None:
                    h1._p.addprevious(h2el)
                    h1._p.addprevious(tbl)

    def _rewrite_text(self, p_el, new_text):
        runs = p_el.findall(qn("w:r"))
        if not runs:
            r = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = new_text
            r.append(t)
            p_el.append(r)
            return
        keep = runs[0]
        for r in runs[1:]:
            p_el.remove(r)
        for t in keep.findall(qn("w:t")):
            keep.remove(t)
        for br in keep.findall(qn("w:br")):
            keep.remove(br)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = new_text
        keep.append(t)

    # ---------- 第一章统计表 ----------
    def update_ch1_tables(self, reg):
        t_year = t_type = None
        for t in self.doc.tables:
            hdr = [c.text.strip() for c in t.rows[0].cells]
            if hdr[:3] == ["年份", "投稿数", "占比"] and t_year is None:
                t_year = t
            elif hdr[:3] == ["类型", "数量", "占比"] and t_type is None:
                t_type = t
        total = len(reg["videos"])
        ycount = Counter(v["date"][:4] for v in reg["videos"] if v.get("date"))
        seen_years = []
        for row in t_year.rows[1:]:
            y = row.cells[0].text.strip()
            seen_years.append(y)
            n = ycount.get(y, 0)
            tcs = row._tr.findall(qn("w:tc"))
            self._cell_set(tcs[0], y)
            self._cell_set(tcs[1], str(n))
            self._cell_set(tcs[2], f"{n / total * 100:.1f}%" if total else "0%")
        for y in sorted(set(ycount) - set(seen_years)):
            n = ycount[y]
            self._table_row_clone(t_year._tbl, [y, str(n), f"{n / total * 100:.1f}%"])
        if t_type is not None:
            tycount = Counter(v["type"] for v in reg["videos"])
            seen_types = []
            for row in t_type.rows[1:]:
                ty = row.cells[0].text.strip()
                seen_types.append(ty)
                n = tycount.get(ty, 0)
                tcs = row._tr.findall(qn("w:tc"))
                self._cell_set(tcs[1], str(n))
                self._cell_set(tcs[2], f"{n / total * 100:.1f}%" if total else "0%")
            for ty in sorted(set(tycount) - set(seen_types)):
                n = tycount[ty]
                extra = ["—"] if len(t_type.rows[0].cells) > 3 else []
                self._table_row_clone(t_type._tbl, [ty, str(n), f"{n / total * 100:.1f}%"] + extra)

    # ---------- 附录 ----------
    def update_appendix(self, reg, new_videos):
        def by_header(sig):
            for t in self.doc.tables:
                hdr = [c.text.strip() for c in t.rows[0].cells]
                if hdr[:len(sig)] == list(sig):
                    yield t

        for t in by_header(["编号", "日期", "标题", "类型", "距上次(天)"]):
            for v in new_videos:
                self._table_row_clone(t._tbl, [str(v["no"]), v["date"].replace("-", "/"), v["title"], v["type"], str(v.get("gap_days") or "首期")])
            break
        vids = [v for v in reg["videos"] if v.get("date")]
        gaps = []
        for i in range(1, len(vids)):
            d0 = datetime.datetime.strptime(vids[i - 1]["date"], "%Y-%m-%d")
            d1 = datetime.datetime.strptime(vids[i]["date"], "%Y-%m-%d")
            gaps.append({"no": vids[i]["no"], "title": vids[i]["title"], "date": vids[i]["date"], "days": (d1 - d0).days})

        def rebuild(t, data, unit):
            rows = t.rows
            tpl_row = copy.deepcopy(rows[-1]._tr)
            for row in rows[1:]:
                t._tbl.remove(row._tr)
            for rank, item in enumerate(data, 1):
                tr = copy.deepcopy(tpl_row)
                tcs = tr.findall(qn("w:tc"))
                vals = [str(rank), str(item["no"]), item["title"], item["date"].replace("-", "/"), f"{item['days']} {unit}"]
                for tc, text in zip(tcs, vals):
                    self._cell_set(tc, text)
                t._tbl.append(tr)

        gap_tabs = list(by_header(["排名", "编号", "标题", "日期", "间隔天数"]))
        if gap_tabs:
            rebuild(gap_tabs[0], sorted(gaps, key=lambda g: -g["days"])[:10], "天")
        for t in by_header(["排名", "编号", "标题", "日期", "字幕行数"]):
            rows = t.rows
            tpl_row = copy.deepcopy(rows[-1]._tr)
            for row in rows[1:]:
                t._tbl.remove(row._tr)
            top = sorted(reg["videos"], key=lambda v: -v.get("sub_lines", 0))[:20]
            for rank, v in enumerate(top, 1):
                tr = copy.deepcopy(tpl_row)
                tcs = tr.findall(qn("w:tc"))
                vals = [str(rank), str(v["no"]), v["title"], (v.get("date") or "").replace("-", "/"), f"{v.get('sub_lines', 0)} 行"]
                for tc, text in zip(tcs, vals):
                    self._cell_set(tc, text)
                t._tbl.append(tr)
            break

    # ---------- 封面 / 统计文本 ----------
    def update_cover(self, reg):
        total = len(reg["videos"])
        main = sum(1 for v in reg["videos"] if v.get("account_short") == "主号")
        alt = total - main
        lines_total = reg.get("total_sub_lines") or sum(v.get("sub_lines", 0) for v in reg["videos"])
        dates = sorted(v["date"] for v in reg["videos"] if v.get("date"))
        repl = [
            (re.compile(r"^收录视频："), f"收录视频：{total} 部（主号 {main} 部 + 小号 {alt} 部）"),
            (re.compile(r"^时间跨度："), f"时间跨度：{dates[0]} 至 {dates[-1]}" if dates else None),
            (re.compile(r"^字幕总行数："), f"字幕总行数：{lines_total:,} 行" if lines_total else None),
        ]
        for p in self.doc.paragraphs:
            t = p.text.strip()
            for pat, new in repl:
                if new and pat.match(t):
                    self._rewrite_text(p._p, new)
                    break
        h1 = self.find_h1("第十二章")
        if h1:
            m = re.match(r"^(第十二章\s*视频词条详录（全)(\d+)(部）)$", h1.text.strip())
            if m:
                self._rewrite_text(h1._p, f"{m.group(1)}{total}{m.group(3)}")

    # ---------- 保存 ----------
    def save_next_version(self):
        m = re.search(r"-v(\d+)\.docx$", self.path.name)
        nxt = int(m.group(1)) + 1 if m else 3
        name = re.sub(r"-v\d+\.docx$", f"-v{nxt}.docx", self.path.name)
        out = self.path.parent / name
        self.doc.save(str(out))
        return str(out)


def build():
    registry = load_json(ROOT / "registry.json")
    videos = registry["videos"]
    meta_cache = load_json(ROOT / "tools" / "bili_meta_cache.json")
    profiles = load_json(ROOT / "tools" / "up_profiles.json")
    insights = load_json(ROOT / "tools" / "danmaku_insights.json")
    account_stats = load_json(ROOT / "tools" / "account_stats.json")
    comments = load_json(ROOT / "tools" / "bili_comments_cache.json")
    for video in videos:
        video["_meta"] = meta_cache.get(video["bvid"], {})

    quiz = parse_quiz()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("怒九百科")
    set_run(run, size=34, bold=True, color=ACCENT)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("怒九笑（主号）× 怒九摸鱼馆（小号）公开作品全量研究")
    set_run(run, size=15, bold=True)
    dates = [v["date"] for v in videos if v.get("date")]
    main_count = sum(v["account_short"] == "主号" for v in videos)
    cover_lines = [
        f"收录视频：{len(videos)} 部（主号 {main_count} 部 + 小号 {len(videos) - main_count} 部）",
        f"时间跨度：{min(dates)} 至 {max(dates)}",
        f"字幕总行数：{registry['total_sub_lines']:,} 行",
        f"弹幕原始记录：{insights.get('total_records', 0):,} 条",
    ]
    for line in cover_lines:
        cover = doc.add_paragraph()
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(cover.add_run(line), size=10.5)
    para(doc, "本文档基于 2026-10-05 抓取的 B 站公开数据生成。视频元数据、字幕、弹幕、评论与账号档案均按可复现口径整理；无 AI/CC 字幕的视频保留公开元数据和补录摘要，不虚构台词。", style="Intense Quote")

    doc.add_page_break()
    heading(doc, "第一章  怒九概览", 1)
    heading(doc, "1.1 基本信息", 2)
    para(doc, "怒九，B站常用名“怒九笑”，是活跃于游戏区、绘画/手书、生活区和音乐区的多栖创作者；其小号“怒九摸鱼馆”承担日常、Vlog、合作和轻松短内容。主号收录 166 部，小号收录 39 部，合计 205 部。")
    para(doc, "她从《守望先锋手书》和 Undertale 相关创作进入收录样本，早期同时尝试手书、游戏推荐、恐怖游戏挑战、恋爱游戏解说与脑洞日常。2018—2020 年形成以像素游戏、独立游戏、恋爱模拟和脑洞动画为核心的创作群；2021 年起与 Warma、捏碳等创作者的联动明显增多；2024—2026 年则进入双人实况、爆米花电台和大型游戏展会记录并行的成熟期。")
    para(doc, "在内容气质上，“怒九”的高频意象是九某人、脑洞日常、艺术就是___、恋爱游戏吐槽、真实反应和强互动合作。她并非单一区创作者：137 部游戏实况构成最大底盘，23 部绘画/手书保留美术功底，13 部搞笑娱乐和 9 部翻唱/音乐补足人设表达。")
    para(doc, "收录口径说明：本百科同时跟踪两个账号的公开投稿。主号“怒九笑”承担正式作品与大型企划，小号“怒九摸鱼馆”承载轻量日常、绘画过程与合作作品；两部分合并编号、统一按投稿时间排序，并保留账号字段以区分来源。任何只看主号或只看小号的统计，都会低估她的实际产出与联动密度。")

    heading(doc, "1.2 创作风格与特点", 2)
    para(doc, "第一，游戏实况占据主体。怒九偏爱独立游戏、像素游戏、恐怖游戏、恋爱模拟与双人合作游戏，玩法解说之外更强调“真实反应”。《双影奇境》达到 791 万播放，是全样本播放量最高的作品，也说明双人实况已成为强牵引力。")
    para(doc, "第二，绘画与手书是身份底色。从《守望先锋手书》《UT手书》到“艺术就是___”系列，她把美术生经验转化为可传播的梗，评论区也经常出现“自学成才”“专业”“美术功底”等反馈。")
    para(doc, "第三，语言风格以短促、直给、吐槽和真实感为主。全站最高频弹幕包括“啊？”“真实”“俺也一样”“好耶”等，说明观众容易把她的反应转译成可复用的共鸣句式。名梗《满糖都是墙》及其变体、“怒九特制麻辣兔头”“谢谢猫猫！猫门永存！”等均形成了稳定的弹幕记忆。")
    para(doc, "第四，合作不是偶发事件。与 Warma 的双人游戏、旅游、电台、搬家、星露谷等内容，与捏碳的短剧和游戏介绍，构成了外部合作网络。小号“怒九摸鱼馆”则承接更轻的日常与摸鱼向作品，避免主号节奏被碎片内容稀释。")
    para(doc, "第五，双账号不是简单分工，而是同一人格的两个出口。主号作品追求完成度，小号作品保留过程感；当主号进入大型双人企划时，小号同时以绘画日常维持更新频率。观众在两个频道之间来回迁移，使梗、称呼和内部笑话得以共享，也让“怒九”作为一个整体形象而非单个频道存在。")

    heading(doc, "1.3 数据统计", 2)
    total_view = sum(stat_of(v, "view") for v in videos)
    total_like = sum(stat_of(v, "like") for v in videos)
    total_danmaku_stat = sum(stat_of(v, "danmaku") for v in videos)
    total_reply_stat = sum(stat_of(v, "reply") for v in videos)
    stats_lines = [
        f"投稿总数：{len(videos)} 部",
        f"时间跨度：{min(dates)} 至 {max(dates)}",
        f"字幕总行数：{registry['total_sub_lines']:,} 行",
        f"字幕总字符：{sum(v['sub_chars'] for v in videos):,} 字符",
        f"有字幕视频：{sum(v['sub_lines'] > 0 for v in videos)} 部",
        f"弹幕原始记录：{insights.get('total_records', 0):,} 条",
        f"视频总播放：{fmt_num(total_view)}",
        f"视频总点赞：{fmt_num(total_like)}",
        f"公开互动总弹幕：{fmt_num(total_danmaku_stat)}",
        f"公开互动总评论：{fmt_num(total_reply_stat)}",
    ]
    for line in stats_lines:
        para(doc, line, bold=line.startswith(("投稿总数", "时间跨度")))
    years = Counter(v["date"][:4] for v in videos)
    add_table(doc, ["年份", "投稿数", "占比"], [[year, years[year], f"{years[year] / len(videos) * 100:.1f}%"] for year in sorted(years)], widths=[1.2, 1.2, 1.2])
    type_count = Counter(v["type"] for v in videos)
    add_table(doc, ["类型", "数量", "占比", "总播放"], [[typ, count, f"{count / len(videos) * 100:.1f}%", fmt_compact(sum(stat_of(v, 'view') for v in videos if v['type'] == typ))] for typ, count in type_count.most_common()], widths=[1.7, 0.8, 0.8, 1.1])
    para(doc, "弹幕与评论口径：弹幕数据合并了最近 XML 与历史分段接口记录，受平台缓存限制，早期视频的弹幕可能不全；评论缓存按点赞排序抓取每支视频的公开热门评论。两者用于刻画互动氛围，不用于绝对值排名。")

    heading(doc, "1.4 数据可视化", 2)
    viz_items = [
        ("01_yearly_trend.png", "年度投稿趋势与累计增长"),
        ("02_type_donut.png", "内容类型分布"),
        ("03_monthly_heatmap.png", "年份 × 月份投稿热力图"),
        ("04_subtitle_growth.png", "字幕规模累计增长"),
        ("05_account_dist.png", "主号与小号投稿分布"),
        ("06_upload_gaps.png", "投稿间隔分析"),
        ("07_top_subtitle.png", "字幕量最多的视频"),
        ("08_intro_freq.png", "自我介绍频率变化"),
        ("09_topics.png", "话题关键词分布"),
        ("10_catchphrases.png", "口头禅与语气词频率"),
        ("11_type_evolution.png", "内容类型年度演变"),
        ("12_type_scatter.png", "类型 × 平均字幕量"),
        ("13_publish_clock.png", "发布时段分布"),
        ("14_danmaku_peak.png", "弹幕峰值窗口分布"),
        ("15_keyword_freq.png", "标题关键词频率"),
        ("16_type_duration.png", "类型 × 平均时长"),
    ]
    for filename, caption in viz_items:
        add_image(doc, filename, f"图：{caption}（2026-10-05 数据）")

    heading(doc, "1.5 标志性口头禅与语言习惯", 2)
    phrase_counter = Counter()
    phrase_videos = defaultdict(set)
    for bvid, item in insights.get("videos", {}).items():
        for phrase in item.get("top_phrases", []):
            phrase_counter[phrase["content"]] += phrase["count"]
            phrase_videos[phrase["content"]].add(bvid)
    add_table(
        doc,
        ["排名", "口头禅 / 高频弹幕", "总重复次数", "出现视频数"],
        [[rank, phrase, count, len(phrase_videos[phrase])] for rank, (phrase, count) in enumerate(phrase_counter.most_common(25), 1)],
        widths=[0.6, 3.2, 1.1, 1.0],
    )
    para(doc, "这里的“口头禅”包含字幕高频表达和弹幕高重复语，两者共同构成粉丝语言的显性层。短句占比高，说明怒九作品更适合即时共鸣式弹幕；“真实”“俺也一样”“好耶”等把个人经验转成公共梗，也解释了评论区的高度互动。")
    para(doc, "弹幕与评论区的分工也很清晰：弹幕负责即时共鸣，短句在特定时间点集中爆发；评论区负责二次创作与提醒，常出现“又来一遍”“考古打卡”和UP主之间的互相留言。理解这一层，才能理解为什么怒九作品的播放未必最高，但互动密度经常领先。")

    heading(doc, "1.6 创作里程碑与粉丝增长", 2)
    milestone_rows = [
        ["2017-08-25", "第 1 部收录作品《守望先锋手书》发布", "第1部"],
        ["2018-01-13", "Undertale 系列开启，游戏与手书双线并进", "第2部"],
        ["2019-03-17", "直播录像类首部收录作品上线", "第44部"],
        ["2021-04-02", "与 Warma 的爆炸电台系列首次收录", "第93部"],
        ["2022-12-16", "“艺术就是___”绘画系列首次出现", "第126部"],
        ["2024-02-24", "《鬼打墙了！！！！》强化双人恐怖实况", "第145部"],
        ["2025-09-10", "爆米花电台第 2 期收录，系列重启", "第182部"],
        ["2026-09-11 / 10-03", "游戏展与出国 Vlog 连发，完成最新阶段记录", "第204-205部"],
    ]
    add_table(doc, ["时间", "事件", "相关视频"], milestone_rows, widths=[1.3, 3.8, 1.3])
    snapshot = account_stats["snapshots"][-1]
    for name, profile in profiles.items():
        fans = snapshot["accounts"][name]["fans"]
        likes = snapshot["accounts"][name]["likes"]
        para(doc, f"{name}：UID {profile['mid']}，粉丝 {fmt_compact(fans)}，获赞 {fmt_compact(likes)}。")
    para(doc, "主号 214.3 万粉丝、小号 60.0 万粉丝，说明两个账号并不是简单的内容备份关系，而是分别承担主创作与轻量陪伴的双频道结构。")

    heading(doc, "1.7 数据口径与更新机制", 2)
    para(doc, "字幕来源：全部字幕来自 B 站 AI/CC 字幕接口，按标点与时间轴合并成段后写入词条；无字幕视频不做台词虚构，只保留公开元数据与补录摘要。")
    para(doc, "弹幕缓存：tools/fetch_danmaku.py 合并实时 XML 与历史分段接口，tools/danmaku_insights.json 输出每支视频的高频弹幕、峰值窗口和年度曲线。")
    para(doc, "评论缓存：tools/bili_comments_cache.json 保存按点赞排序的热门评论及部分楼中楼回复，供词条与网站引用。")
    para(doc, "更新机制：tools/update.py 负责检查新投稿、下载字幕并调用 build_docx.DocPatcher 增量写入词条、年表、统计表与封面；全量重建走 build_docx.py，每次输出新版本号。")

    doc.add_page_break()
    heading(doc, "第二章  创作年表", 1)
    by_year = defaultdict(list)
    for video in videos:
        by_year[video["date"][:4]].append(video)
    year_order = sorted(by_year)
    for year in year_order:
        rows = by_year[year]
        heading(doc, f"2.{year_order.index(year) + 1} {year}年：{len(rows)}部投稿", 2)
        year_views = sum(stat_of(v, "view") for v in rows)
        year_likes = sum(stat_of(v, "like") for v in rows)
        year_types = Counter(v["type"] for v in rows).most_common(3)
        para(doc, f"年度规模：{len(rows)} 部；播放 {fmt_compact(year_views)}；点赞 {fmt_compact(year_likes)}；主要类型：{'、'.join(f'{k}({v})' for k, v in year_types)}。")
        add_table(
            doc,
            ["编号", "日期", "标题", "类型"],
            [[v["no"], v["date"], v["title"], v["type"]] for v in rows],
            widths=[0.5, 0.9, 4.2, 0.8],
        )

    doc.add_page_break()
    heading(doc, "第三章  内容类型分析", 1)
    heading(doc, "3.1 类型演变与体量对比", 2)
    para(doc, "怒九的类型结构呈“游戏实况为底盘、绘画/手书为身份、生活与合作为外延”的三层结构。游戏实况占 66.8%，但 2021 年后双人合作和电台内容占比上升；绘画/手书从 2017—2022 年稳定输出，2024 年后减少，与其转向双人实况和旅游记录同步。")
    para(doc, "与单区创作者不同，怒九的类型迁移有清晰的因果链：游戏实况积累反应素材，绘画/手书沉淀美术身份，两者在合作与电台中合流为“真实人格输出”。因此本章既按类型分节，也保留跨类型的联动观察。")
    add_image(doc, "11_type_evolution.png", "图：内容类型年度演变")
    add_image(doc, "12_type_scatter.png", "图：类型 × 平均字幕量对比")
    type_notes = {
        "游戏实况": "从早期单机实况到双人合作，反应型解说贯穿始终，玩法之外更看人。",
        "绘画/手书": "美术生出身的底色，使手书成为频道最具辨识度的体裁。",
        "搞笑娱乐": "以脑洞短剧和挑战类内容为主，承担破圈与引流功能。",
        "翻唱/音乐": "唱歌多与剧情、绘画叙事结合，少有孤立的技术型翻唱。",
        "知识科普": "用轻松包装讨论偏见、教育与社会观察，扩展频道表达半径。",
        "配音/小剧场": "角色化配音放大性格差异，是小剧场喜剧的主要载体。",
        "爆炸电台": "与好友的长谈形式，把生活事件放进固定栏目。",
        "生活日常": "主号与小号交替记录日常，维持账号的陪伴感。",
        "直播录像": "直播切片沉淀为存档，保留最即时的互动状态。",
    }
    for section_index, typ in enumerate(sorted(type_count, key=lambda x: type_count[x], reverse=True), 2):
        rows = [v for v in videos if v["type"] == typ]
        heading(doc, f"3.{section_index} {typ}", 2)
        avg_duration = statistics.mean(v["duration"] for v in rows)
        avg_lines = statistics.mean(v["sub_lines"] for v in rows)
        para(doc, f"共 {len(rows)} 部，平均时长 {fmt_duration(avg_duration)}，平均字幕 {avg_lines:.0f} 行，总播放 {fmt_compact(sum(stat_of(v, 'view') for v in rows))}。")
        year_dist = Counter(v["date"][:4] for v in rows)
        para(doc, "时间分布：" + "、".join(f"{y} 年 {n} 部" for y, n in sorted(year_dist.items())) + f"。最早一期是第 {rows[0]['no']} 部《{rows[0]['title']}》，最新一期是第 {rows[-1]['no']} 部《{rows[-1]['title']}》。")
        para(doc, "体裁解读：" + type_notes.get(typ, "该类型与频道其他体裁交叉明显，常以不同形式复用同一创意。"))
        top_rows = sorted(rows, key=lambda v: stat_of(v, "view"), reverse=True)[:8]
        add_table(
            doc,
            ["编号", "日期", "标题", "播放", "点赞"],
            [[v["no"], v["date"], v["title"], fmt_compact(stat_of(v, "view")), fmt_compact(stat_of(v, "like"))] for v in top_rows],
            widths=[0.45, 0.85, 3.75, 0.75, 0.75],
        )

    doc.add_page_break()
    heading(doc, "第四章  爆炸电台 / 爆米花电台专题", 1)
    radio = [v for v in videos if v["type"] == "爆炸电台"]
    heading(doc, "4.1 系列概述", 2)
    para(doc, "怒九收录样本中的电台作品共 4 部，跨越 2021—2026 年。系列前段以“姐妹俩打打闹闹的日常”出现，2025 年转入“爆米花电台”编号，2026 年又补充游戏展与出国记录。它既是与 Warma 的合作记录，也是把旅行、受伤恢复、游戏展等生活事件放进长谈的形式实验。")
    heading(doc, "4.2 各期一览", 2)
    add_table(
        doc,
        ["编号", "标题", "日期", "播放", "点赞", "时长"],
        [[v["no"], v["title"], v["date"], fmt_compact(stat_of(v, "view")), fmt_compact(stat_of(v, "like")), fmt_duration(v["duration"])] for v in radio],
        widths=[0.45, 3.55, 0.95, 0.7, 0.7, 0.65],
    )
    heading(doc, "4.3 各期解析", 2)
    for video in radio:
        heading(doc, f"第{video['no']}部：{video['title']}", 3)
        para(doc, f"发布于 {video['date']}，时长 {fmt_duration(video['duration'])}，播放 {fmt_compact(stat_of(video, 'view'))}，点赞 {fmt_compact(stat_of(video, 'like'))}。该期属于 {video['type']}，账号为 {video['account']}。")
        para(doc, fallback_summary(video))
        for line in fallback_highlights(video):
            para(doc, line, style="List Bullet")
        add_comments(doc, video, comments, 2)
        add_danmaku(doc, video, insights)
        radio_item = insights.get("videos", {}).get(video["bvid"])
        if radio_item and radio_item.get("top_phrases"):
            para(doc, "弹幕侧写：观众用「" + "、".join(x["content"] for x in radio_item["top_phrases"][:4]) + "」等短语实时接话，高互动来自共同话题而非单向收听。")

    heading(doc, "4.4 系列演变观察", 2)
    para(doc, "2021 年的爆炸电台是“顺带发生”的：两人在游戏实况间隙闲聊，被完整记录下来，成为系列的原点。它证明这对搭档即使不设计环节，也能靠日常对话撑起一整期内容。")
    para(doc, "2025 年重启为“爆米花电台”后，系列进入编号化阶段：固定开场、固定互相吐槽、固定生活近况汇报，形式感明显增强。编号化同时意味着档期化——观众开始期待下一期，而不是把电台当成实况的附属品。")
    para(doc, "2026 年的几期把电台带出室内：游戏展见闻、出国记录被放进长谈框架里复盘。电台由此从“聊天栏目”升级为“生活事件的消化现场”，这也是它持续有生命力的原因。")

    doc.add_page_break()
    heading(doc, "第五章  生活/Vlog 系列专题", 1)
    life_pat = r"Vlog|旅游|出国|搬家|日常"
    life = [v for v in videos if v["type"] == "生活日常" or re.search(life_pat, v["title"], re.I)]
    heading(doc, "5.1 系列概述", 2)
    para(doc, f"生活向内容共 {len(life)} 部（含“生活日常”类型与标题含 Vlog/旅游/出国/搬家/日常的作品）。它不追求完成度，而是把频道从“作品集合”还原成“人的记录”：搬家、旅行、出国、摸鱼日常，构成了词条之间最松弛的过渡带。")
    para(doc, "生活系列的载体横跨两个账号：主号负责旅行与展会级别的记录，小号负责更私人的日常碎片。这种分布与“正式作品在主号、陪伴内容在小号”的双频道战略一致。")
    heading(doc, "5.2 全量作品一览", 2)
    add_table(
        doc,
        ["编号", "日期", "标题", "账号", "类型", "播放"],
        [[v["no"], v["date"], v["title"], v["account_short"], v["type"], fmt_compact(stat_of(v, "view"))] for v in life],
        widths=[0.45, 0.85, 3.2, 0.55, 0.75, 0.7],
    )
    heading(doc, "5.3 阶段演变", 2)
    para(doc, "早期（2018—2020）的生活向内容以“日常”为题眼的短记录为主，功能是调剂游戏实况的节奏；搬家与旅行主题在 2021 年后明显增多，与合作关系深化同步。")
    para(doc, "2024—2026 年，生活系列进入“事件级”阶段：出国看展、游戏展见闻被当作小型企划制作，标题不再只是“日常”，而是带有地点与事件的具体名词。生活内容由此从碎片升级为可期待的系列。")
    para(doc, "值得注意的是，生活系列的平均字幕量偏低，但弹幕互动密度不低——观众对“人的状态”比对“作品质量”更容易产生即时共鸣，这也是双账号结构能维持粘性的关键。")

    doc.add_page_break()
    heading(doc, "第六章  游戏实况专题", 1)
    games = [v for v in videos if v["type"] == "游戏实况"]
    heading(doc, "6.1 游戏内容结构", 2)
    para(doc, "游戏实况共 137 部，占全部投稿 66.8%，累计播放 1.83 亿、点赞 994.3 万，平均时长 30 分 59 秒。它不是单一路径，而是由独立像素推荐、恐怖挑战、恋爱模拟、硬核动作、双人合作、平台解谜和模拟经营共同组成。")
    add_table(
        doc,
        ["题材", "代表视频数", "首次代表作品", "作用"],
        [
            ["Undertale / Deltarune", "5+", "第2部《Undertale：某地下的御茶会议》", "早期人设与手书联动的核心 IP"],
            ["像素 / 独立游戏推荐", "9+", "第9部《五款神作级别的像素游戏》", "建立速推荐与宝藏游戏口碑"],
            ["恐怖游戏", "14+", "第14部《恐怖游戏大挑战》", "真实惊叫和吐槽反应的素材库"],
            ["恋爱模拟 / 撩到我算我输", "15+", "第40部《情人节特别篇》", "把乙女游戏做成吐槽连续剧"],
            ["双人合作", "30+", "第93部与 Warma 的电台实况", "2024—2026 强增长板块"],
            ["Splatoon / 动作竞技", "5+", "第123部《4k狙也太难用了吧》", "展现“菜但敢玩”的竞技人格"],
        ],
        widths=[1.5, 0.8, 2.5, 1.6],
    )
    heading(doc, "6.2 重点游戏系列", 2)
    series_patterns = {
        "Undertale / Deltarune": r"undertale|deltarune|ut手书|ut人类组|传说之下|三角符文",
        "双影奇境": r"双影奇境",
        "星露谷物语": r"星露谷|田园开荒",
        "Splatoon": r"splatoon",
        "撩到我算我输": r"撩到我|恋爱游戏|乙女游戏",
        "国产像素 / 独立游戏": r"国产良心像素|像素游戏|国产游戏|独立游戏",
    }
    for series, pattern in series_patterns.items():
        rows = [v for v in games if re.search(pattern, v["title"] + " " + " ".join(v.get("tags", [])), re.I)]
        if rows:
            para(doc, f"{series}：{len(rows)} 部，首次出现在第 {rows[0]['no']} 部《{rows[0]['title']}》。", style="List Bullet")
    heading(doc, "6.3 高播放游戏实况", 2)
    add_table(
        doc,
        ["排名", "标题", "日期", "播放", "点赞"],
        [[rank, v["title"], v["date"], fmt_compact(stat_of(v, "view")), fmt_compact(stat_of(v, "like"))] for rank, v in enumerate(sorted(games, key=lambda x: stat_of(x, "view"), reverse=True)[:20], 1)],
        widths=[0.45, 4.1, 0.9, 0.7, 0.7],
    )

    doc.add_page_break()
    heading(doc, "第七章  音乐作品专题", 1)
    music = [v for v in videos if v["type"] == "翻唱/音乐"]
    para(doc, f"音乐类共 {len(music)} 部，平均时长 {fmt_duration(statistics.mean(v['duration'] for v in music))}，总播放 {fmt_compact(sum(stat_of(v, 'view') for v in music))}。这些作品较少走纯技术翻唱路线，更多把唱歌、叙事动画、兄妹梗和日常剧情结合在一起。")
    add_table(
        doc,
        ["编号", "标题", "日期", "播放", "点赞", "字幕"],
        [[v["no"], v["title"], v["date"], fmt_compact(stat_of(v, "view")), fmt_compact(stat_of(v, "like")), v["sub_lines"]] for v in music],
        widths=[0.45, 3.9, 0.9, 0.7, 0.7, 0.65],
    )
    para(doc, "重点包括《我画了一本书！再不进来听就变成黑历史了！！》的绘画 + 声音叙事，《【warma/怒九】让我们快乐地搬家吧！》的合作合唱感，以及“有个亲哥”系列中的角色化配音表达。音乐不是孤立技能展示，而是脑洞内容的延长线。")

    heading(doc, "第八章  绘画/手书 与 知识科普/配音专题", 1)
    art = [v for v in videos if v["type"] == "绘画/手书"]
    para(doc, f"绘画/手书共 {len(art)} 部，总播放 {fmt_compact(sum(stat_of(v, 'view') for v in art))}。样本从第 1 部《守望先锋手书》开始，经 Undertale、求生之路、V家手书，到 2022 年后的“艺术就是___”系列，形成美术生身份的连续证明。")
    heading(doc, "8.1 绘画/手书全量作品表", 2)
    add_table(
        doc,
        ["编号", "标题", "日期", "播放", "点赞", "字幕"],
        [[v["no"], v["title"], v["date"], fmt_compact(stat_of(v, "view")), fmt_compact(stat_of(v, "like")), v["sub_lines"]] for v in art],
        widths=[0.45, 3.9, 0.9, 0.7, 0.7, 0.65],
    )
    para(doc, "“艺术就是___”系列的重要价值在于把创作过程本身娱乐化：她让观众看错误、改稿、失控和童画直觉，而不是只展示成品。这与绘画区常见的技巧展示形成差异，也更贴近她的“陪伴感”人格。")

    heading(doc, "8.2 知识科普与配音/小剧场", 2)
    special = [v for v in videos if v["type"] in ("知识科普", "配音/小剧场")]
    para(doc, f"知识科普 {sum(v['type'] == '知识科普' for v in videos)} 部，配音/小剧场 {sum(v['type'] == '配音/小剧场' for v in videos)} 部。两者数量不大，但承担了讨论偏见、教育经历、社会观察和角色扮演的功能，让频道不只是游戏记录。")
    add_table(
        doc,
        ["编号", "类型", "标题", "日期", "播放", "点赞"],
        [[v["no"], v["type"], v["title"], v["date"], fmt_compact(stat_of(v, "view")), fmt_compact(stat_of(v, "like"))] for v in special],
        widths=[0.45, 0.9, 3.5, 0.85, 0.65, 0.65],
    )
    para(doc, "《艺术生遭受到了哪些偏见？》《LGBT群体遭到了哪些误解？》《全国统一的人类共同行为记录》等作品显示，她能以轻松包装处理社会观察和身份议题；配音短剧则通过夸张情境放大性格差异。")

    doc.add_page_break()
    heading(doc, "第九章  合作关系", 1)
    heading(doc, "9.1 与 Warma 的合作", 2)
    warma_rows = [v for v in videos if re.search(r"warma|沃玛", v["title"] + " " + " ".join(v.get("tags", [])), re.I)]
    para(doc, f"标题或标签中明确关联 Warma / 沃玛的作品共 {len(warma_rows)} 部，覆盖双人游戏、旅游、电台、搬家、赏月、生日蛋糕和游戏展等场景。两人合作不是简单互客，而是“反应互补”式联动：Warma 提供稳定叙事，怒九提供爆发吐槽。")
    add_table(
        doc,
        ["编号", "日期", "标题", "类型", "播放"],
        [[v["no"], v["date"], v["title"], v["type"], fmt_compact(stat_of(v, "view"))] for v in warma_rows[:40]],
        widths=[0.45, 0.85, 3.9, 0.85, 0.75],
    )
    heading(doc, "9.2 与捏碳及其他合作对象", 2)
    partner_patterns = {
        "捏碳": r"捏碳",
        "四迹": r"四迹",
        "衔九 / 衔生": r"衔九|衔生",
        "亲哥 / 家庭成员": r"亲哥|哥哥|兄妹",
    }
    for partner, pattern in partner_patterns.items():
        rows = [v for v in videos if re.search(pattern, v["title"] + " " + " ".join(v.get("tags", [])), re.I)]
        if rows:
            para(doc, f"{partner}：关联 {len(rows)} 部，代表作 {rows[0]['title']}。", style="List Bullet")
    para(doc, "合作网络还包含评论区与联名梗。Warma、捏碳碳碳碳、四迹、warma养鸽场等账号都出现在高赞评论中，说明怒九的内容生态和 Warma 圈层高度互通，同时保留自己的“真实吐槽”声音。")

    heading(doc, "第十章  怒九熟悉度测试", 1)
    quiz_categories = sorted(set(item["category"] for item in quiz))
    para(doc, f"本章收录 {len(quiz)} 道自动生成题目，覆盖 {len(quiz_categories)} 个分类和简单 / 中等 / 困难三档。题目来自 registry、meta、标签、弹幕、评论和账号档案，可用来检验读者对数据细节的掌握。")
    quiz_by_category = defaultdict(list)
    for item in quiz:
        quiz_by_category[item["category"]].append(item)
    category_no = 0
    for category in quiz_categories:
        rows = quiz_by_category.get(category)
        if not rows:
            continue
        category_no += 1
        heading(doc, f"10.{category_no} {category}", 2)
        diff_map = {"easy": "简单", "medium": "中等", "hard": "困难"}
        table_rows = []
        for idx, item in enumerate(rows, 1):
            table_rows.append([idx, diff_map.get(item["difficulty"], item["difficulty"]), item["q"], item["options"][item["answer"]], item.get("explanation", "")])
        add_table(doc, ["#", "难度", "问题", "正确答案", "解析"], table_rows, widths=[0.35, 0.6, 2.25, 1.4, 1.7])

    doc.add_page_break()
    heading(doc, "第十一章  小号视频专题（怒九摸鱼馆）", 1)
    alt = [v for v in videos if v["account_short"] == "小号"]
    heading(doc, "11.1 账号定位", 2)
    para(doc, f"小号收录 {len(alt)} 部，粉丝 {fmt_compact(snapshot['accounts']['小号（怒九摸鱼馆）']['fans'])}。它更像“轻松创作间”，内容以绘画、翻唱、日常、游戏碎片和合作短作品为主，节奏比主号更松散。")
    alt_type = Counter(v["type"] for v in alt)
    add_table(doc, ["类型", "数量", "总播放"], [[typ, count, fmt_compact(sum(stat_of(v, 'view') for v in alt if v['type'] == typ))] for typ, count in alt_type.most_common()], widths=[1.6, 0.8, 1.1])
    heading(doc, "11.2 小号全量列表", 2)
    add_table(
        doc,
        ["编号", "日期", "标题", "类型", "播放"],
        [[v["no"], v["date"], v["title"], v["type"], fmt_compact(stat_of(v, "view"))] for v in alt],
        widths=[0.45, 0.85, 3.9, 0.85, 0.75],
    )

    doc.add_page_break()
    heading(doc, f"第十二章  视频词条详录（全{len(videos)}部）", 1)
    para(doc, "本章按投稿时间升序收录全部作品。有字幕视频写入关键话题、口头禅、摘要、亮点、氛围、语言特征和合并字幕；无字幕视频提供公开数据补录摘要、标签、互动和弹幕特征。词条中所有数字均为 2026-10-05 快照。")
    for video in videos:
        add_video_entry(doc, video, comments, insights)

    doc.add_page_break()
    heading(doc, "附录  统计数据汇总", 1)
    heading(doc, "A.1 投稿时间线", 2)
    add_table(
        doc,
        ["编号", "日期", "标题", "类型", "距上次(天)"],
        [[v["no"], v["date"], v["title"], v["type"], v.get("gap_days") or "首期投稿"] for v in videos],
        widths=[0.45, 0.85, 3.85, 0.85, 0.8],
    )
    gap_rows = [v for v in videos if isinstance(v.get("gap_days"), int)]
    heading(doc, "A.2 间隔 TOP10", 2)
    add_table(
        doc,
        ["排名", "编号", "标题", "日期", "间隔天数"],
        [[rank, v["no"], v["title"], v["date"], v["gap_days"]] for rank, v in enumerate(sorted(gap_rows, key=lambda x: x["gap_days"], reverse=True)[:10], 1)],
        widths=[0.5, 0.5, 3.7, 0.9, 0.7],
    )
    heading(doc, "A.3 字幕行数 TOP20", 2)
    add_table(
        doc,
        ["排名", "编号", "标题", "日期", "字幕行数"],
        [[rank, v["no"], v["title"], v["date"], v["sub_lines"]] for rank, v in enumerate(sorted(videos, key=lambda x: x["sub_lines"], reverse=True)[:20], 1)],
        widths=[0.5, 0.5, 3.7, 0.9, 0.7],
    )
    heading(doc, "A.4 播放量 TOP20", 2)
    add_table(
        doc,
        ["排名", "编号", "标题", "日期", "播放"],
        [[rank, v["no"], v["title"], v["date"], fmt_compact(stat_of(v, "view"))] for rank, v in enumerate(sorted(videos, key=lambda x: stat_of(x, "view"), reverse=True)[:20], 1)],
        widths=[0.5, 0.5, 3.7, 0.9, 0.7],
    )
    heading(doc, "A.5 点赞量 TOP20", 2)
    add_table(
        doc,
        ["排名", "编号", "标题", "日期", "点赞"],
        [[rank, v["no"], v["title"], v["date"], fmt_compact(stat_of(v, "like"))] for rank, v in enumerate(sorted(videos, key=lambda x: stat_of(x, "like"), reverse=True)[:20], 1)],
        widths=[0.5, 0.5, 3.7, 0.9, 0.7],
    )
    peak_rows = []
    for video in videos:
        item = insights.get("videos", {}).get(video["bvid"])
        if item and item.get("peaks"):
            peak_rows.append((video, item["peaks"][0]))
    heading(doc, "A.6 弹幕峰值 TOP20", 2)
    add_table(
        doc,
        ["排名", "编号", "标题", "峰值时间", "峰值条数"],
        [[rank, v["no"], v["title"], fmt_duration(p.get("t", 0)), p.get("count", 0)] for rank, (v, p) in enumerate(sorted(peak_rows, key=lambda x: x[1].get("count", 0), reverse=True)[:20], 1)],
        widths=[0.5, 0.5, 3.5, 0.85, 0.8],
    )
    heading(doc, "A.7 可视化索引", 2)
    add_table(
        doc,
        ["图号", "文件", "主题"],
        [[idx + 1, filename, caption] for idx, (filename, caption) in enumerate(viz_items)],
        widths=[0.5, 2.2, 3.6],
    )

    versions = [int(m.group(1)) for p in ROOT.glob("nujiu-encyclopedia-v*.docx") for m in [re.search(r"-v(\d+)\.docx$", p.name)] if m]
    out = ROOT / f"nujiu-encyclopedia-v{max(versions, default=1) + 1}.docx"
    doc.save(out)
    check = Document(str(out))
    print(f"saved {out} | paras={len(check.paragraphs)} tables={len(check.tables)} | {out.stat().st_size / 1024 / 1024:.1f} MB")


if __name__ == "__main__":
    build()
